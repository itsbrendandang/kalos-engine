"""Review fixes for #87's upload path: string column names, xlsx scan bounds
that do not trust a sheet's declared dimension, too-few-rows before
"constant", dates refused as targets, and clipped / unlogged reflected names.
All workbooks are synthetic and built in memory."""
from __future__ import annotations

import io
import logging
import time
import zipfile

import numpy as np
import openpyxl
import pandas as pd
import pytest

from kalos.portal import uploads
from kalos.portal.uploads import (
    MAX_REFLECTED_CHARS,
    MAX_REFLECTED_NAMES,
    UploadRejected,
    clip_name,
    parse_upload,
    safe_context,
)
from kalos.validation.target import assess_target_variation

pytest.importorskip("fastapi")
from fastapi.testclient import TestClient  # noqa: E402

from kalos.portal import app as portal_module  # noqa: E402
from kalos.portal import campaign as campaign_module  # noqa: E402
from kalos.portal.analysis import _analyze  # noqa: E402
from kalos.portal.campaign import CampaignStore, get_campaign_store  # noqa: E402


@pytest.fixture
def client(tmp_path, monkeypatch):
    store = CampaignStore(tmp_path)
    monkeypatch.setattr(portal_module, "_STATE_DIR", tmp_path)
    monkeypatch.setattr(portal_module, "_LATEST_DIR", tmp_path / "latest")
    monkeypatch.setattr(campaign_module, "_STORE", store)
    portal_module.app.dependency_overrides[get_campaign_store] = lambda: store
    try:
        yield TestClient(portal_module.app)
    finally:
        portal_module.app.dependency_overrides.pop(get_campaign_store, None)


@pytest.fixture
def no_fit(monkeypatch):
    import kalos.core.optimize as optimize
    import kalos.core.surrogate as surrogate

    def _boom(*_a, **_k):
        raise AssertionError("a model was fit / a batch proposed")

    monkeypatch.setattr(surrogate.Surrogate, "fit", _boom)
    monkeypatch.setattr(optimize, "propose_batch", _boom)


def _sheet(n: int = 20, seed: int = 0) -> pd.DataFrame:
    rng = np.random.default_rng(seed)
    m = rng.uniform(0, 4, n)
    return pd.DataFrame({
        "Methanol": m.round(3),
        "pH": rng.uniform(5, 7, n).round(2),
        "lipase_titer": (1.5 * m + rng.normal(0, 0.2, n)).round(3),
    })


# --- (H) column names are always strings ----------------------------------------- #


def _xlsx_with_int_header() -> bytes:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["Methanol", "pH", 1])  # the target's header cell is the NUMBER 1
    df = _sheet()
    for row in df.itertuples(index=False):
        ws.append(list(row))
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def test_numeric_header_cells_become_string_column_names():
    parsed = parse_upload(_xlsx_with_int_header())
    assert list(parsed.df.columns) == ["Methanol", "pH", "1"]
    assert all(isinstance(c, str) for c in parsed.df.columns)


def test_csv_column_names_are_strings_too():
    parsed = parse_upload(b"1,2,3\n" + b"4,5,6\n" * 8)
    assert list(parsed.df.columns) == ["1", "2", "3"]


def test_a_numeric_header_can_be_requested_as_the_target(client):
    r = client.post(
        "/api/run",
        files={"file": ("runs.xlsx", _xlsx_with_int_header(), "application/octet-stream")},
        data={"target": "1"},
    )
    assert r.status_code == 200, r.text
    assert r.json()["target"] == "1"


# --- (I) xlsx scan bounds that do not trust the declared dimension --------------- #

_NS = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"


def _sheet_xml(rows: int, *, dimension: str | None, broken: bool = False) -> bytes:
    head = '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
    head += f'<worksheet xmlns="{_NS}">'
    if dimension:
        head += f'<dimension ref="{dimension}"/>'
    body = ["<sheetData>"]
    body.append('<row r="1"><c r="A1" t="inlineStr"><is><t>a</t></is></c>'
                '<c r="B1" t="inlineStr"><is><t>b</t></is></c></row>')
    for r in range(2, rows + 2):
        body.append(f'<row r="{r}"><c r="A{r}"><v>{r}</v></c><c r="B{r}"><v>{r * 0.5}</v></c></row>')
    tail = "<row><c r=" if broken else "</sheetData></worksheet>"  # broken: truncated XML
    return (head + "".join(body) + tail).encode()


def _workbook(n_sheets: int) -> bytes:
    wb = openpyxl.Workbook()
    wb.active.title = "s1"
    for i in range(2, n_sheets + 1):
        wb.create_sheet(f"s{i}")
    for ws in wb.worksheets:
        ws["A1"] = "x"
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _replace_sheets(raw: bytes, parts: dict[str, bytes]) -> bytes:
    src = zipfile.ZipFile(io.BytesIO(raw))
    out = io.BytesIO()
    with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as dst:
        for info in src.infolist():
            dst.writestr(info.filename, parts.get(info.filename, src.read(info.filename)))
    return out.getvalue()


def test_lying_dimension_sheet_is_refused_before_parsing(monkeypatch):
    """A sheet that declares A1:B2 but holds far more is refused from the zip's
    uncompressed sizes, before pandas/openpyxl parse anything."""
    xml = _sheet_xml(40_000, dimension="A1:B2")
    assert len(xml) > 1_000_000
    raw = _replace_sheets(_workbook(1), {"xl/worksheets/sheet1.xml": xml})
    monkeypatch.setattr(uploads, "MAX_XLSX_UNCOMPRESSED_BYTES", 1_000_000)

    def _never(*_a, **_k):
        raise AssertionError("the workbook was parsed")

    monkeypatch.setattr(uploads.pd, "ExcelFile", _never)
    with pytest.raises(UploadRejected, match="too large"):
        parse_upload(raw)


def test_lying_dimension_under_the_byte_cap_is_still_caught_after_parse(monkeypatch):
    """Belt and braces: under the byte cap, the materialized frame is checked."""
    xml = _sheet_xml(3_000, dimension="A1:B2")
    raw = _replace_sheets(_workbook(1), {"xl/worksheets/sheet1.xml": xml})
    monkeypatch.setattr(uploads, "MAX_XLSX_CELLS", 1_000)
    with pytest.raises(UploadRejected, match="too large"):
        parse_upload(raw)


def test_uncompressed_cap_default_is_64_mib():
    assert uploads.MAX_XLSX_UNCOMPRESSED_BYTES == 64 * 1024 * 1024


def test_twenty_broken_sheets_finish_quickly_and_are_charged():
    """20 sheets that each fail at the very end of a ~1.8 MB part: each failure
    is charged its part's cell bound, so the scan stops early instead of
    parsing all of them, and the request finishes in a few seconds."""
    parts = {
        f"xl/worksheets/sheet{i}.xml": _sheet_xml(25_000, dimension="A1:B2", broken=True)
        for i in range(1, 21)
    }
    raw = _replace_sheets(_workbook(20), parts)
    total = sum(len(p) for p in parts.values())
    assert total < uploads.MAX_XLSX_UNCOMPRESSED_BYTES  # the byte cap alone would allow it

    calls: list[str] = []
    real_parse = pd.ExcelFile.parse

    def _counting(self, name, *a, **k):
        calls.append(name)
        return real_parse(self, name, *a, **k)

    with pytest.MonkeyPatch.context() as mp:
        mp.setattr(pd.ExcelFile, "parse", _counting)
        t0 = time.perf_counter()
        with pytest.raises(UploadRejected) as exc:
            parse_upload(raw)
        elapsed = time.perf_counter() - t0
    assert exc.value.code == "no_usable_sheet"
    # Each failure is charged part_size // 16 cells; the scan stops once the
    # MAX_XLSX_CELLS budget is spent (it would try all 20 if failures were free).
    per_sheet = len(parts["xl/worksheets/sheet1.xml"]) // uploads._MIN_CELL_XML_BYTES
    expected = -(-uploads.MAX_XLSX_CELLS // per_sheet)
    assert expected < 20
    assert len(calls) == expected, calls
    assert elapsed < 10, elapsed


def test_a_broken_sheet_does_not_block_a_good_one_behind_it():
    good = _sheet()
    buf = io.BytesIO()
    with pd.ExcelWriter(buf) as w:
        pd.DataFrame({"x": [1]}).to_excel(w, sheet_name="broken", index=False)
        good.to_excel(w, sheet_name="data", index=False)
    raw = _replace_sheets(buf.getvalue(), {"xl/worksheets/sheet1.xml": _sheet_xml(10, dimension="A1:B11", broken=True)})
    parsed = parse_upload(raw)
    assert parsed.sheet_used == "data"


# --- (J) too few rows before "constant"; dates are not numeric ------------------- #


def test_few_identical_target_values_is_too_few_rows_not_constant(no_fit):
    df = _sheet()
    df["lipase_titer"] = [5.0, 5.0, 5.0] + [np.nan] * (len(df) - 3)
    with pytest.raises(UploadRejected) as exc:
        _analyze(df, "lipase_titer")
    assert exc.value.code == "too_few_rows"
    assert exc.value.context["n_rows"] == 3
    assert "same value" not in exc.value.message


def test_enough_identical_target_values_is_still_constant(no_fit):
    df = _sheet()
    df["lipase_titer"] = 5.0
    with pytest.raises(UploadRejected) as exc:
        _analyze(df, "lipase_titer")
    assert exc.value.code == "target_constant"


@pytest.mark.parametrize("values", [
    pd.Series(pd.date_range("2024-01-01", periods=10)),
    pd.Series(pd.date_range("2024-01-01", periods=10, tz="UTC")),
    pd.Series(pd.to_timedelta(np.arange(10), unit="h")),
])
def test_datetime_and_timedelta_targets_are_not_numeric(values):
    verdict = assess_target_variation(values)
    assert verdict.status == "not_numeric" and verdict.n_numeric == 0


def test_a_date_target_is_refused_by_analyze(no_fit):
    df = _sheet()
    df["harvest_date"] = pd.date_range("2024-01-01", periods=len(df))
    with pytest.raises(UploadRejected) as exc:
        _analyze(df, "harvest_date")
    assert exc.value.code == "target_not_numeric"


# --- (K) reflected names are clipped and capped; logs carry codes only ----------- #


def test_clip_name_and_safe_context():
    long = "x" * 500
    assert len(clip_name(long)) == MAX_REFLECTED_CHARS
    ctx = safe_context({"target": long, "available_targets": [f"c{i}" for i in range(50)], "n_rows": 3})
    assert len(ctx["target"]) == MAX_REFLECTED_CHARS
    assert len(ctx["available_targets"]) == MAX_REFLECTED_NAMES
    assert ctx["available_targets_total"] == 50
    assert ctx["n_rows"] == 3


def test_long_requested_target_is_clipped_in_the_error(client, caplog):
    secret_long = "SECRETCOL" + "y" * 400
    raw = _sheet().to_csv(index=False).encode()
    with caplog.at_level(logging.WARNING, logger="kalos.portal"):
        r = client.post("/api/run", files={"file": ("runs.csv", raw, "text/csv")},
                        data={"target": secret_long})
    assert r.status_code == 400
    body = r.json()
    assert body["code"] == "target_not_found"
    assert len(body["target"]) == MAX_REFLECTED_CHARS
    assert secret_long not in body["error"]
    assert len(body["error"]) < 400
    # the server log names the code, never the customer's column names
    assert "code=target_not_found" in caplog.text
    assert "SECRETCOL" not in caplog.text and "Methanol" not in caplog.text


def test_wide_sheet_available_targets_are_capped(client):
    rng = np.random.default_rng(0)
    df = pd.DataFrame({f"titer_{i}": rng.uniform(0, 1, 10) for i in range(30)})
    raw = df.to_csv(index=False).encode()
    r = client.post("/api/run", files={"file": ("runs.csv", raw, "text/csv")}, data={"target": "nope"})
    body = r.json()
    assert body["code"] == "target_not_found"
    assert len(body["available_targets"]) == MAX_REFLECTED_NAMES
    assert body["available_targets_total"] == 30


def test_many_sheet_names_are_capped_in_the_message():
    sheets = {f"notes{i}": pd.DataFrame({"a": [1]}) for i in range(30)}
    buf = io.BytesIO()
    with pd.ExcelWriter(buf) as w:
        for name, df in sheets.items():
            df.to_excel(w, sheet_name=name, index=False)
    with pytest.raises(UploadRejected) as exc:
        parse_upload(buf.getvalue())
    body = exc.value.to_body()
    assert body["code"] == "no_usable_sheet"
    assert "and 10 more" in body["error"] and "notes29" not in body["error"]
    assert len(body["sheets"]) == MAX_REFLECTED_NAMES and body["sheets_total"] == 30


def test_anonymize_pseudonymizes_identifier_names_in_the_error(client):
    df = _sheet()
    df.insert(0, "batch_id", np.arange(len(df)) + 100)
    df = df.rename(columns={"lipase_titer": "y"})
    raw = df.to_csv(index=False).encode()
    plain = client.post("/api/run", files={"file": ("runs.csv", raw, "text/csv")},
                        data={"target": "nope"}).json()
    assert "batch_id" in plain["available_targets"]
    anon = client.post("/api/run", files={"file": ("runs.csv", raw, "text/csv")},
                       data={"target": "nope", "anonymize": "true"}).json()
    assert anon["code"] == "target_not_found"
    assert "batch_id" not in anon["available_targets"]
    assert any(n.startswith("col_") for n in anon["available_targets"])
    assert "Methanol" in anon["available_targets"]  # process inputs keep their names
