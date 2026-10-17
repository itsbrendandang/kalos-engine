"""Kalos portal — the untrusted-input boundary.

This engine ingests untrusted run sheets from external clients, so the raw
upload and its parsed shape are capped to bound memory and blunt zip-bomb
expansion. Sizes are configurable via env; the CSV/xlsx caps are constants.
"""
from __future__ import annotations

import io
import logging
import os
import re
import zipfile
from dataclasses import dataclass, field

import pandas as pd

from kalos.normalize.orientation import detect_orientation

log = logging.getLogger("kalos.portal")

# --- upload safety limits ---------------------------------------------------- #
# `or` (not the .get default) so a SET-BUT-EMPTY var falls back too: compose
# and k8s templates pass `KALOS_MAX_UPLOAD_MB=${KALOS_MAX_UPLOAD_MB:-}`, which
# arrives as "" - .get's default never applies and float("") crashed the boot
# (found on the first real container run). Same unset==empty convention
# kalos/providers/ already follows for every credential var.
_MAX_UPLOAD_MB = float(os.environ.get("KALOS_MAX_UPLOAD_MB", "").strip() or "25")
MAX_UPLOAD_BYTES = int(_MAX_UPLOAD_MB * 1024 * 1024)
MAX_CSV_ROWS = 100_000       # rows read from a CSV/TSV upload
MAX_COLUMNS = 512            # columns allowed in any upload (CSV or xlsx)
MAX_XLSX_CELLS = 2_000_000   # rows * cols ceiling for a parsed xlsx (zip-bomb guard)
MAX_XLSX_UNCOMPRESSED_BYTES = 64 * 1024 * 1024
# Ceiling on the TOTAL uncompressed size of every part of an xlsx, from the zip
# central directory, checked BEFORE anything is parsed. Python's zipfile never
# yields more than an entry's declared `file_size` (it truncates and then fails
# the CRC), so this bounds every byte openpyxl/pandas can read out of the
# workbook - whatever a sheet's `<dimension>` claims - and therefore the total
# parse work across all sheets (~5 s of openpyxl at ~13 MB/s). 64 MiB is about
# 1.5M simple numeric cells, far above any run sheet the GP can fit
# (MAX_FIT_ROWS x MAX_COLUMNS = ~1M cells); in practice it binds before
# MAX_XLSX_CELLS.
_MIN_CELL_XML_BYTES = 16
# A sheet cell costs at least this many bytes of XML (`<c r="A1"><v>1</v></c>`
# is 22), so `part_size // _MIN_CELL_XML_BYTES` is an upper bound on the cells a
# sheet part can hold. Used to charge the default-sheet scan budget for a sheet
# that failed mid-parse, whose real cell count is unknown.
_ZIP_MAGIC = b"PK\x03\x04"   # xlsx/xls-as-zip start-of-file marker

# Separate, tighter cap on the rows the SURROGATE is actually fit on. A
# SingleTaskGP is O(n^2) in memory and O(n^3) in time, so a raw upload that is
# within MAX_CSV_ROWS can still be far too large for an exact GP. The optimizer
# targets the small-sample bioprocess regime (N <= a couple thousand), so we
# reject an over-cap fit set rather than silently subsampling (which would be
# invisible, non-deterministic data loss). Override with KALOS_MAX_FIT_ROWS.
MAX_FIT_ROWS = int(os.environ.get("KALOS_MAX_FIT_ROWS", "").strip() or "2000")

MIN_FIT_ROWS = 6
# The fewest data rows (with a numeric target) a model is fit on - the floor
# `kalos.portal.analysis._analyze` has always enforced, named here so the upload
# error can state it.

# Non-leaking messages. We never echo the parser/exception text or a cell value.
# Each message below is ENGINE-AUTHORED and names the specific problem and what
# to do about it; the only user content that may appear in one is a column or
# sheet name the uploader themselves supplied (the success response returns
# those names anyway).
_ERR_PARSE = "Could not parse the uploaded file. Check it is a CSV or Excel run-sheet."
_ERR_EMPTY = (
    "The uploaded file is empty. Upload a CSV or Excel (.xlsx) run sheet with a "
    "header row and at least one data row per run."
)
_ERR_UNSUPPORTED = (
    "Unsupported file type{what}. Upload the run sheet as CSV or tab-separated text "
    "(.csv, .tsv, .txt) or as an Excel workbook (.xlsx)."
)
_ERR_LEGACY_EXCEL = (
    "This looks like a legacy Excel 97-2003 (.xls) or password-protected workbook, "
    "which cannot be read. Open it in Excel and save it as an unprotected .xlsx "
    "workbook, or export the sheet as CSV."
)
_ERR_CORRUPT_EXCEL = (
    "The Excel file could not be opened: it looks corrupt or is not a valid .xlsx "
    "workbook. Re-save it from Excel as .xlsx, or export the run sheet as CSV."
)
_ERR_NOT_A_TABLE = (
    "The file could not be read as a table: its lines do not have a consistent "
    "number of comma- or tab-separated fields. Export the run sheet from Excel as "
    "CSV, or upload the .xlsx workbook."
)
_ERR_TOO_LARGE = "The uploaded file is too large."
_ERR_TOO_MANY_COLUMNS = "The uploaded file has too many columns."
_ERR_TOO_MANY_FIT_ROWS = (
    "The dataset is too large for the surrogate; the optimizer targets the "
    f"small-sample regime, N<={MAX_FIT_ROWS}."
)
# A numerically-hard-but-valid file (near-duplicate or ill-conditioned rows) that
# defeats even the jittered fit retry gets its OWN message, so it is not confused
# with the generic parse failure.
_ERR_FIT = (
    "The model could not be fit on this data - likely near-duplicate or "
    "ill-conditioned rows."
)


MAX_REFLECTED_CHARS = 100
MAX_REFLECTED_NAMES = 20
# Error responses (and their messages) may name a column or sheet the uploader
# supplied. Such reflected strings are clipped to MAX_REFLECTED_CHARS and name
# lists capped at MAX_REFLECTED_NAMES (`<key>_total` then carries the true
# count), so a crafted header cannot turn an error body into a large or
# arbitrary-content echo.


def clip_name(name: object) -> str:
    """`str(name)`, clipped to `MAX_REFLECTED_CHARS` (with a trailing ellipsis)."""
    text = str(name)
    if len(text) <= MAX_REFLECTED_CHARS:
        return text
    return text[: MAX_REFLECTED_CHARS - 1] + "\u2026"


def clip_names(names) -> list[str]:
    """At most `MAX_REFLECTED_NAMES` names, each clipped (see `clip_name`)."""
    return [clip_name(n) for n in list(names)[:MAX_REFLECTED_NAMES]]


def safe_context(context: dict) -> dict:
    """An error context with every reflected string clipped and every list of
    names capped; a capped list adds `<key>_total` with its full length."""
    out: dict = {}
    for key, value in context.items():
        if isinstance(value, str):
            out[key] = clip_name(value)
        elif isinstance(value, (list, tuple)) and all(isinstance(v, str) for v in value):
            out[key] = clip_names(value)
            if len(value) > MAX_REFLECTED_NAMES:
                out[f"{key}_total"] = len(value)
        else:
            out[key] = value
    return out


def scrubbed_failure(exc: BaseException) -> str:
    """A log-safe description of an unexpected failure: the exception TYPE and
    where it was raised (the innermost few `file:line in function` frames),
    never `str(exc)`. pandas/openpyxl/numpy messages routinely quote cell
    values and column names ("could not convert string to float: 'ACME-...'"),
    and those must not reach the server log; the type plus location is enough
    to find the bug."""
    import traceback

    frames = traceback.extract_tb(exc.__traceback__)[-3:]
    where = " <- ".join(
        f"{os.path.basename(f.filename)}:{f.lineno} in {f.name}" for f in reversed(frames)
    )
    return f"{type(exc).__name__}" + (f" at {where}" if where else "")


def rejection_counts(rej: "UploadRejected") -> str:
    """A log suffix with only the numeric facts of a rejection (`n_rows=4
    min_rows=6`, list lengths as `n_<key>`), never names or values, so the
    server log carries no customer column or sheet names."""
    parts: list[str] = []
    for key, value in rej.context.items():
        if isinstance(value, bool):
            continue
        if isinstance(value, int):
            parts.append(f"{key}={value}")
        elif isinstance(value, (list, tuple)):
            parts.append(f"n_{key}={len(value)}")
    return (" " + " ".join(parts)) if parts else ""


class UploadRejected(ValueError):
    """A client upload failed a safety guard or cannot be analyzed.

    Carries a safe, engine-authored message (never parser/exception text), plus
    an optional machine-readable `code` and a small `context` dict of extra
    JSON-safe fields (e.g. `{"target": ..., "n_rows": ...}`). `to_body()` is
    the `{"error": ..., "code": ..., **context}` envelope every route returns,
    so an old client that only reads `error` keeps working unchanged.
    """

    def __init__(
        self, message: str, *, code: str | None = None, context: dict | None = None
    ) -> None:
        super().__init__(message)
        self.code = code
        self.context: dict = dict(context or {})

    @property
    def message(self) -> str:
        return str(self.args[0]) if self.args else ""

    def to_body(self) -> dict:
        body: dict = {"error": self.message}
        if self.code:
            body["code"] = self.code
        for key, value in safe_context(self.context).items():
            if key not in ("error", "code"):
                body[key] = value
        return body


_SHEET_XML_RE = re.compile(r"^xl/worksheets/sheet\d+\.xml$")
_DIMENSION_RE = re.compile(rb'<dimension\s+ref="([^"]+)"\s*/?>')
_CELL_REF_RE = re.compile(r"^([A-Za-z]+)(\d+)$")
# The <dimension> tag is schema-guaranteed to precede the (potentially huge)
# <sheetData> section, so a bounded prefix read is enough to find it - this
# keeps the guard cheap even against a sheet whose real (not just declared)
# content is itself a compressed-XML bomb.
_DIMENSION_SCAN_BYTES = 65_536


def _col_letters_to_index(letters: str) -> int:
    """Convert a 1-based spreadsheet column label (A, Z, AA, SF, ...) to an index."""
    idx = 0
    for ch in letters.upper():
        idx = idx * 26 + (ord(ch) - ord("A") + 1)
    return idx


def _parse_dimension_ref(ref: str) -> tuple[int, int] | None:
    """Parse a `<dimension ref="...">` value (e.g. `A1:SF1048576`, or a single
    cell like `A1`) into its 1-based (rows, cols) extent. Returns None if the
    bottom-right cell reference does not match the expected `<letters><digits>`
    shape (defensive - callers fall back to openpyxl in that case).
    """
    end = ref.split(":")[-1].strip()
    m = _CELL_REF_RE.match(end)
    if not m:
        return None
    col_letters, row_str = m.groups()
    return int(row_str), _col_letters_to_index(col_letters)


_SPREADSHEETML_NS = frozenset({
    "http://schemas.openxmlformats.org/spreadsheetml/2006/main",
    "http://purl.oclc.org/ooxml/spreadsheetml/main",  # ISO strict
})
_EXTENT_SCAN_CHUNK = 64 * 1024


class _NotASheet(Exception):
    """Stop scanning a zip part whose root element is not a worksheet."""


class _SheetExtentScan:
    """Streaming (expat) scan of one sheet part for the extent pandas will
    materialize: the largest row index and column index of any `<row>` /
    `<c>` inside `<sheetData>`, resolved the way openpyxl does (an explicit
    `r` attribute, else one past the previous row / cell). Raises
    `UploadRejected` the moment the extent passes the column or cell cap, so a
    sparse sheet (one cell far to the right on every row, no or a lying
    `<dimension>`) is refused after reading a few KB, never expanded.
    Memory is O(1): no tree is built."""

    def __init__(self) -> None:
        self.max_row = 0
        self.max_col = 0
        self._row = 0
        self._col = 0
        self._in_data = 0
        self._seen_root = False

    def _check(self) -> None:
        if self.max_col > MAX_COLUMNS:
            raise UploadRejected(_ERR_TOO_MANY_COLUMNS)
        if self.max_row * max(self.max_col, 1) > MAX_XLSX_CELLS:
            raise UploadRejected(_ERR_TOO_LARGE)

    def start(self, name: str, attrs: dict[str, str]) -> None:
        ns, _, local = name.rpartition(" ")
        if not self._seen_root:
            self._seen_root = True
            if local not in ("worksheet", "dialogsheet"):
                raise _NotASheet
        if ns not in _SPREADSHEETML_NS:
            return
        if local == "sheetData":
            self._in_data += 1
        elif self._in_data and local == "row":
            r = attrs.get("r", "")
            self._row = int(r) if r.isdigit() else self._row + 1
            self._col = 0
            self.max_row = max(self.max_row, self._row)
            self._check()
        elif self._in_data and local == "c":
            m = _CELL_REF_RE.match(attrs.get("r", ""))
            if m:
                self._col = _col_letters_to_index(m.group(1))
                self.max_row = max(self.max_row, int(m.group(2)))
            else:
                self._col += 1
            self.max_col = max(self.max_col, self._col)
            self._check()

    def end(self, name: str) -> None:
        if name.rpartition(" ")[2] == "sheetData" and self._in_data:
            self._in_data -= 1


def _scan_sheet_extents(zf: zipfile.ZipFile) -> None:
    """Run `_SheetExtentScan` over EVERY zip part whose root element is a
    worksheet (whatever the part is named, since openpyxl follows the
    workbook's relationships, not a naming convention). Total work is bounded
    by `MAX_XLSX_UNCOMPRESSED_BYTES`; a part that is not XML, or not a sheet,
    stops after its first chunk. A sheet whose XML breaks part-way keeps the
    extent read so far (openpyxl would have read that much too) and is left
    for the parser to report as corrupt."""
    import pyexpat

    for info in zf.infolist():
        if info.is_dir():
            continue
        scan = _SheetExtentScan()
        parser = pyexpat.ParserCreate(namespace_separator=" ")
        parser.SetParamEntityParsing(pyexpat.XML_PARAM_ENTITY_PARSING_NEVER)
        parser.StartElementHandler = scan.start
        parser.EndElementHandler = scan.end
        try:
            with zf.open(info) as part:
                while chunk := part.read(_EXTENT_SCAN_CHUNK):
                    parser.Parse(chunk, False)
                parser.Parse(b"", True)
        except (_NotASheet, pyexpat.ExpatError):
            continue
        except UploadRejected:
            log.warning("rejected an xlsx upload: a sheet's real extent is over the cap")
            raise
        except (zipfile.BadZipFile, OSError, EOFError, ValueError) as exc:
            log.warning("rejected an unreadable xlsx upload: %s", type(exc).__name__)
            raise UploadRejected(_ERR_CORRUPT_EXCEL, code="corrupt_excel") from exc


def _reject_oversized_xlsx_via_openpyxl(raw: bytes) -> None:
    """Fallback guard for the rare sheet with no `<dimension>` tag: open read-only
    with openpyxl and check its computed `max_row` / `max_column` instead.
    """
    import openpyxl

    try:
        wb = openpyxl.load_workbook(io.BytesIO(raw), read_only=True)
    except Exception as exc:  # noqa: BLE001 - corrupt/non-xlsx zip -> specific safe 400
        log.warning("rejected an unreadable xlsx upload: %s", type(exc).__name__)
        raise UploadRejected(_ERR_CORRUPT_EXCEL, code="corrupt_excel") from exc
    try:
        for ws in wb.worksheets:
            if ws.max_column is None or ws.max_row is None:
                # Unknown extent (no dimension): never read as 0. The streaming
                # extent scan (`_scan_sheet_extents`) has already bounded it.
                continue
            cols = ws.max_column
            rows = ws.max_row
            if cols > MAX_COLUMNS:
                raise UploadRejected(_ERR_TOO_MANY_COLUMNS)
            if rows * cols > MAX_XLSX_CELLS:
                raise UploadRejected(_ERR_TOO_LARGE)
    finally:
        wb.close()


def _reject_oversized_xlsx(raw: bytes) -> dict[str, int]:
    """Reject an xlsx whose declared dimensions exceed the cell / column caps,
    or whose parts add up to more than `MAX_XLSX_UNCOMPRESSED_BYTES`
    uncompressed, BEFORE `pd.read_excel` materializes the frame (the zip-bomb
    guard). Returns `{part name: uncompressed size}` for every zip entry.

    The uncompressed-size total is the guard that does not trust the sheet: a
    `<dimension>` tag can understate the real content (a "lying" sheet), but
    the bytes behind it are capped by the zip's declared sizes, which zipfile
    enforces on read.

    Reads each sheet's `<dimension ref="...">` tag straight out of the raw zip
    entry (`xl/worksheets/sheet*.xml`) via a lightweight regex - this is the
    DECLARED shape a crafted workbook can inflate independently of its real
    content, so it must be checked without ever asking openpyxl or pandas to
    walk (and thus materialize) the sheet. If a sheet has no `<dimension>` tag,
    falls back to opening it read-only with openpyxl and using its computed
    `max_row` / `max_column`. A corrupt or unreadable zip raises
    `UploadRejected` (client error).
    """
    try:
        zf = zipfile.ZipFile(io.BytesIO(raw))
    except zipfile.BadZipFile as exc:
        log.warning("rejected an unreadable xlsx upload: %s", type(exc).__name__)
        raise UploadRejected(_ERR_CORRUPT_EXCEL, code="corrupt_excel") from exc

    with zf:
        try:
            names = set(zf.namelist())
        except Exception as exc:  # noqa: BLE001 - a damaged central directory
            raise UploadRejected(_ERR_CORRUPT_EXCEL, code="corrupt_excel") from exc
        if "xl/workbook.xml" not in names:
            # A well-formed zip that is not an Excel workbook: a Word or
            # PowerPoint file, or a plain archive. Say what it is, not "corrupt".
            if any(n.startswith("word/") for n in names):
                what = " (this is a Word document)"
            elif any(n.startswith("ppt/") for n in names):
                what = " (this is a PowerPoint file)"
            elif any(n.startswith("xl/") for n in names):
                raise UploadRejected(_ERR_CORRUPT_EXCEL, code="corrupt_excel")
            else:
                what = " (this is a zip archive, not an Excel workbook)"
            raise UploadRejected(_ERR_UNSUPPORTED.format(what=what), code="unsupported_file_type")
        sheet_names = [n for n in zf.namelist() if _SHEET_XML_RE.match(n)]
        part_sizes = {info.filename: int(info.file_size) for info in zf.infolist()}
        if sum(part_sizes.values()) > MAX_XLSX_UNCOMPRESSED_BYTES:
            log.warning("rejected an xlsx upload over the uncompressed-size cap")
            raise UploadRejected(_ERR_TOO_LARGE)
        # The authoritative extent check: stream every sheet part and refuse
        # one whose real (not declared) rows x columns pass the caps, before
        # pandas/openpyxl materialize anything. A declared <dimension> can be
        # missing or lie; the cells' own coordinates cannot be skipped.
        _scan_sheet_extents(zf)
        if not sheet_names:
            # Not a recognizable xlsx layout (or a non-xlsx zip) - let openpyxl's
            # own validation produce the generic parse-rejection.
            _reject_oversized_xlsx_via_openpyxl(raw)
            return part_sizes

        needs_fallback = False
        for name in sheet_names:
            with zf.open(name) as sheet_file:
                head = sheet_file.read(_DIMENSION_SCAN_BYTES)
            match = _DIMENSION_RE.search(head)
            if not match:
                needs_fallback = True
                continue
            parsed = _parse_dimension_ref(match.group(1).decode("ascii", errors="replace"))
            if parsed is None:
                needs_fallback = True
                continue
            rows, cols = parsed
            if cols > MAX_COLUMNS:
                raise UploadRejected(_ERR_TOO_MANY_COLUMNS)
            if rows * cols > MAX_XLSX_CELLS:
                raise UploadRejected(_ERR_TOO_LARGE)

    if needs_fallback:
        _reject_oversized_xlsx_via_openpyxl(raw)
    return part_sizes


def _apply_orientation_prepass(df: pd.DataFrame) -> pd.DataFrame:
    """Tier-1 deterministic pre-pass (`kalos.normalize.orientation`): detect
    whether the just-parsed sheet is standard or transposed orientation, and
    normalize a confidently-transposed sheet BEFORE anything downstream
    (unit normalization, the validation gate, analysis) ever sees it - see
    `kalos/normalize/orientation.py`'s module docstring for why this has to
    happen first.

    Never silently swaps the frame without a trace: the `OrientationReport`
    is always recorded on the returned frame's `df.attrs["kalos_orientation"]`
    - including the "standard"/"ambiguous" no-op cases - so a downstream
    provenance report can see what this pre-pass decided and why, without
    changing `_parse_upload`'s return type or every caller's signature.
    An ambiguous sheet is returned completely untouched (never guessed);
    a standard sheet is also returned untouched, so this pre-pass leaves
    kalos's existing behavior byte-identical for every sheet that was
    already standard orientation.
    """
    report = detect_orientation(df)
    out = report.normalized_frame if report.orientation == "transposed" and report.normalized_frame is not None else df
    out.attrs["kalos_orientation"] = {
        "orientation": report.orientation,
        "confidence": report.confidence,
        "signals": dict(report.signals),
    }
    return out


# --- file-type sniffing ----------------------------------------------------- #

_OLE2_MAGIC = b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1"  # legacy .xls / encrypted OOXML
_BINARY_MAGICS: tuple[tuple[bytes, str], ...] = (
    (b"%PDF", "a PDF"),
    (b"\x89PNG", "a PNG image"),
    (b"\xff\xd8\xff", "a JPEG image"),
    (b"GIF8", "a GIF image"),
    (b"\x1f\x8b", "a gzip archive"),
    (b"7z\xbc\xaf", "a 7z archive"),
    (b"Rar!", "a RAR archive"),
    (b"{\\rtf", "a rich-text document"),
)
# Extensions that are never a run sheet, rejected by name even when the bytes
# would decode as text (a .json or .html file "parses" into nonsense columns).
# A real xlsx is still accepted whatever it is named: magic bytes win when they
# positively identify a workbook (`tests/test_hardening.py` pins that an xlsx
# uploaded as `runs.csv` works).
_NON_TABULAR_EXTS = frozenset({
    ".pdf", ".doc", ".docx", ".ppt", ".pptx", ".odt", ".rtf", ".png", ".jpg",
    ".jpeg", ".gif", ".zip", ".gz", ".7z", ".rar", ".json", ".xml", ".html",
    ".htm", ".md", ".py", ".js", ".yaml", ".yml",
})


def _extension(filename: str | None) -> str:
    return os.path.splitext(filename or "")[1].lower()


def _reject_non_tabular(raw: bytes, filename: str | None) -> None:
    """Refuse bytes that are recognisably not a CSV/TSV/xlsx run sheet."""
    if raw[:8] == _OLE2_MAGIC:
        raise UploadRejected(_ERR_LEGACY_EXCEL, code="unsupported_file_type")
    for magic, what in _BINARY_MAGICS:
        if raw.startswith(magic):
            raise UploadRejected(
                _ERR_UNSUPPORTED.format(what=f" (this looks like {what})"),
                code="unsupported_file_type",
            )
    ext = _extension(filename)
    if ext in _NON_TABULAR_EXTS:
        raise UploadRejected(
            _ERR_UNSUPPORTED.format(what=f" ({ext} files are not run sheets)"),
            code="unsupported_file_type",
        )
    if ext == ".xls":
        # Text bytes under a legacy-Excel name (e.g. an HTML "xls" export).
        raise UploadRejected(_ERR_LEGACY_EXCEL, code="unsupported_file_type")


def _decode_text(raw: bytes) -> str:
    """Decode a text upload. UTF-16 is honoured only via its BOM (Excel's
    "Unicode Text" export); otherwise UTF-8 (a leading BOM is stripped), with
    undecodable bytes replaced as before. A NUL byte in UTF-8 text means the
    file is binary, not a table."""
    if raw[:2] in (b"\xff\xfe", b"\xfe\xff"):
        try:
            return raw.decode("utf-16")
        except UnicodeError as exc:
            raise UploadRejected(_ERR_NOT_A_TABLE, code="not_a_table") from exc
    if b"\x00" in raw[:8192]:
        raise UploadRejected(
            _ERR_UNSUPPORTED.format(what=" (the file is binary, not text)"),
            code="unsupported_file_type",
        )
    return raw.decode("utf-8-sig", errors="replace")


# --- what makes a sheet usable ---------------------------------------------- #


def numeric_columns(df: pd.DataFrame) -> list:
    """Columns where >=80% of the NON-BLANK cells parse as numbers (and there are
    at least 3 of them). Blanks count as "absent", so a sparse component column
    still counts. The single definition `kalos.portal.analysis` also uses."""
    out = []
    for c in df.columns:
        s = df[c]
        if isinstance(s, pd.DataFrame):  # duplicate header; analysis dedupes later
            s = s.iloc[:, 0]
        nonblank = s.notna() & (s.astype(str).str.strip() != "")
        if nonblank.sum() < 3:
            continue
        if pd.to_numeric(s[nonblank], errors="coerce").notna().mean() >= 0.8:
            out.append(c)
    return out


def table_shape_problem(df: pd.DataFrame) -> str | None:
    """Why `df` cannot be a run-sheet table, or None if it plausibly can.

    A NECESSARY condition only (never stricter than what `_analyze` needs): at
    least 2 non-empty columns, at least `MIN_FIT_ROWS` non-empty data rows, and
    at least one numeric column. Used to pick a workbook's default sheet.
    """
    trimmed = df.dropna(axis=1, how="all").dropna(axis=0, how="all")
    if trimmed.shape[1] < 2:
        return "too_few_columns"
    if trimmed.shape[0] < MIN_FIT_ROWS:
        return "too_few_rows"
    if not numeric_columns(trimmed):
        return "no_numeric_inputs"
    return None


# --- parsing ------------------------------------------------------------------ #


MAX_SHEETS_SCANNED = 20
# The default-sheet search parses sheets in workbook order until one is usable.
# Capped (and also bounded by the cumulative MAX_XLSX_CELLS budget below) so a
# workbook of many near-cap sheets cannot multiply the parse cost.


@dataclass
class ParsedUpload:
    """A parsed upload plus where in the file it came from.

    `kind` is "excel" or "text". For Excel, `sheets` lists every sheet name in
    workbook order and `sheet_used` names the one that was parsed (None for a
    text upload, where both are absent from the response). `sheet_requested`
    is True when the caller chose the sheet explicitly.
    """

    df: pd.DataFrame
    kind: str
    sheets: list[str] = field(default_factory=list)
    sheet_used: str | None = None
    sheet_requested: bool = False

    def context(self) -> dict:
        """The keys the API response carries (Excel uploads only)."""
        if self.kind != "excel":
            return {}
        return {"sheets": list(self.sheets), "sheet_used": self.sheet_used}

    def annotate(self, rej: UploadRejected) -> UploadRejected:
        """Return `rej` with this workbook's sheet context, and, for a
        multi-sheet workbook, a sentence naming the sheet that was analyzed and
        the others the caller could choose instead."""
        if self.kind != "excel" or "sheets" in rej.context:
            return rej
        message = rej.message
        others = [n for n in self.sheets if n != self.sheet_used]
        if others and self.sheet_used is not None:
            message += (
                f" (This was sheet {clip_name(self.sheet_used)!r}; the workbook also has "
                f"{_quote_list(others)}. Choose another sheet to analyze it instead.)"
            )
        return UploadRejected(
            message, code=rej.code, context={**rej.context, **self.context()}
        )


def _quote_list(names: list[str]) -> str:
    """Names for a message: clipped, at most `MAX_REFLECTED_NAMES`, then a count."""
    shown = ", ".join(repr(n) for n in clip_names(names))
    extra = len(names) - MAX_REFLECTED_NAMES
    return shown + (f" and {extra} more" if extra > 0 else "")


def _check_frame_caps(df: pd.DataFrame) -> None:
    rows, cols = df.shape
    if cols > MAX_COLUMNS:
        raise UploadRejected(_ERR_TOO_MANY_COLUMNS)
    if rows * cols > MAX_XLSX_CELLS:
        raise UploadRejected(_ERR_TOO_LARGE)


def _match_sheet(requested: str, names: list[str]) -> str | None:
    """Exact sheet name, else a unique case-/whitespace-insensitive match."""
    if requested in names:
        return requested
    key = requested.strip().casefold()
    hits = [n for n in names if n.strip().casefold() == key]
    return hits[0] if len(hits) == 1 else None


def _parse_excel(raw: bytes, sheet: str | None) -> ParsedUpload:
    """Parse one sheet of an xlsx workbook.

    SHEET RULE. `sheet` given: that sheet (exact name, else a unique
    case-insensitive match), or a `sheet_not_found` error listing the sheets.
    `sheet` omitted: the FIRST sheet in workbook order that holds a usable
    table (`table_shape_problem` is None) - for a single-sheet workbook, or one
    whose first sheet is the data, that is exactly the old "first sheet"
    behaviour. If no sheet is usable: a single-sheet workbook is still parsed
    (so the analysis can say precisely what is wrong with it); a multi-sheet
    workbook is refused with `no_usable_sheet` and the sheet list.
    """
    # xlsx/xls: enforce the cell-count ceiling BEFORE pd.read_excel fully
    # materializes the frame, so a zip-bomb whose DECLARED sheet dimensions are
    # enormous is rejected without the memory spike of building the DataFrame.
    part_sizes = _reject_oversized_xlsx(raw)
    try:
        book = pd.ExcelFile(io.BytesIO(raw))
        names = [str(n) for n in book.sheet_names]
    except Exception as exc:  # noqa: BLE001 - any reader failure is a corrupt workbook
        log.warning("rejected an unreadable xlsx upload: %s", type(exc).__name__)
        raise UploadRejected(_ERR_CORRUPT_EXCEL, code="corrupt_excel") from exc
    if not names:
        raise UploadRejected(_ERR_CORRUPT_EXCEL, code="corrupt_excel")
    ctx = {"sheets": names, "sheet_used": None}

    def _read(name: str) -> pd.DataFrame:
        try:
            df = book.parse(name)
        except Exception as exc:  # noqa: BLE001 - e.g. a chart sheet, a damaged part
            log.warning("could not read an xlsx sheet: %s", type(exc).__name__)
            raise UploadRejected(
                f"Sheet {clip_name(name)!r} could not be read. " + _ERR_CORRUPT_EXCEL,
                code="corrupt_excel",
                context={**ctx, "sheet": name},
            ) from exc
        # Belt-and-suspenders: re-check the materialized shape. The pre-read guard
        # uses the sheet's declared dimensions; this catches a mismatch and the
        # column ceiling on the actual parsed frame.
        _check_frame_caps(df)
        return _apply_orientation_prepass(df)

    try:
        if sheet is not None and sheet.strip():
            chosen = _match_sheet(sheet, names)
            if chosen is None:
                raise UploadRejected(
                    f"The workbook has no sheet named {clip_name(sheet)!r}. Its sheets are: "
                    f"{_quote_list(names)}.",
                    code="sheet_not_found",
                    context={**ctx, "sheet": sheet},
                )
            return ParsedUpload(_read(chosen), "excel", names, chosen, sheet_requested=True)

        if len(names) == 1:
            return ParsedUpload(_read(names[0]), "excel", names, names[0])
        # Scan budget in cells, shared by every sheet tried: a parsed sheet is
        # charged its real size; one that failed mid-parse is charged the most
        # cells its XML part could hold (it may have been read almost to the
        # end before failing), so broken sheets cannot multiply the work.
        budget = MAX_XLSX_CELLS
        for name in names[:MAX_SHEETS_SCANNED]:
            try:
                df = _read(name)
            except UploadRejected as rej:
                if rej.code != "corrupt_excel":
                    raise
                # an unreadable sheet is not a candidate; charge it, try the next
                budget -= _sheet_cell_bound(book, name, part_sizes)
                if budget <= 0:
                    break
                continue
            if table_shape_problem(df) is None:
                return ParsedUpload(df, "excel", names, name)
            budget -= int(df.shape[0] * df.shape[1])
            if budget <= 0:
                break
        raise UploadRejected(
            f"None of the workbook's {len(names)} sheets ({_quote_list(names)}) holds a "
            f"usable run-sheet table: one needs a header row, at least 2 columns, at "
            f"least {MIN_FIT_ROWS} data rows, and numeric columns. Put the run sheet "
            "on its own sheet, or export it as CSV.",
            code="no_usable_sheet",
            context=ctx,
        )
    finally:
        book.close()


def _sheet_cell_bound(book: pd.ExcelFile, name: str, part_sizes: dict[str, int]) -> int:
    """Upper bound on the cells sheet `name` can hold, from its XML part's
    uncompressed size. Falls back to an even share of the cell cap when the
    part cannot be located (a non-openpyxl reader, a private attribute gone)."""
    fallback = max(MAX_XLSX_CELLS // MAX_SHEETS_SCANNED, 1)
    try:
        path = getattr(book.book[name], "_worksheet_path", None)
    except Exception:  # noqa: BLE001 - only an estimate; never fail the upload here
        return fallback
    size = part_sizes.get(str(path or "").lstrip("/"))
    return max(size // _MIN_CELL_XML_BYTES, 1) if size else fallback


def _parse_text(raw: bytes) -> ParsedUpload:
    text = _decode_text(raw)
    if text.strip() == "":
        raise UploadRejected(_ERR_EMPTY, code="empty_file")
    head = text[:4000]
    sep = "\t" if head.count("\t") > head.count(",") else ","
    try:
        # cap columns first (cheap, from the header) before reading the full body
        ncols = pd.read_csv(io.StringIO(text), sep=sep, nrows=0).shape[1]
        if ncols > MAX_COLUMNS:
            raise UploadRejected(_ERR_TOO_MANY_COLUMNS)
        df = pd.read_csv(io.StringIO(text), sep=sep)
    except pd.errors.ParserError as exc:
        log.warning("rejected an unparseable text upload: %s", type(exc).__name__)
        raise UploadRejected(_ERR_NOT_A_TABLE, code="not_a_table") from exc
    except pd.errors.EmptyDataError as exc:
        raise UploadRejected(_ERR_EMPTY, code="empty_file") from exc
    # Fail closed on the row cap rather than silently truncating to the first
    # MAX_CSV_ROWS rows: silent data loss would contradict the safe-errors
    # contract, so an over-cap CSV is rejected like the xlsx cell-cap guard.
    if len(df) > MAX_CSV_ROWS:
        raise UploadRejected(_ERR_TOO_LARGE)
    return ParsedUpload(_apply_orientation_prepass(df), "text")


def parse_upload(
    raw: bytes, *, filename: str | None = None, sheet: str | None = None
) -> ParsedUpload:
    """Turn raw upload bytes into a bounded dataframe (plus sheet metadata), or
    raise `UploadRejected` with a specific, safe message and `code`.

    Guards, in order:
      1. byte-size cap (default 25 MB, env `KALOS_MAX_UPLOAD_MB`);
      2. empty file (`empty_file`);
      3. filetype sniff by MAGIC BYTES first: a zip header (`PK\\x03\\x04`) is an
         xlsx candidate (a non-workbook zip such as .docx is
         `unsupported_file_type`, a damaged one `corrupt_excel`); a legacy/
         encrypted OLE2 file, a PDF/image/archive, a binary blob, or a
         non-tabular extension (.pdf, .json, ...) is `unsupported_file_type`;
         anything else is read as UTF-8 (or BOM-marked UTF-16) CSV/TSV, and
         text that is not a consistent table is `not_a_table`;
      4. shape caps: CSV/TSV rows and a column ceiling, and an xlsx cell-count
         (rows * cols) ceiling to blunt zip-bomb expansion.
    `sheet` selects a workbook sheet by name (ignored for text uploads); see
    `_parse_excel` for the default-sheet rule. No message ever contains parser
    or exception text.
    """
    if len(raw) > MAX_UPLOAD_BYTES:
        raise UploadRejected(_ERR_TOO_LARGE)
    if len(raw) == 0 or raw.strip(b" \t\r\n\x00\xef\xbb\xbf\xff\xfe") == b"":
        raise UploadRejected(_ERR_EMPTY, code="empty_file")
    if raw[:4] == _ZIP_MAGIC:
        parsed = _parse_excel(raw, sheet)
    else:
        _reject_non_tabular(raw, filename)
        parsed = _parse_text(raw)
    # Column names are always strings: an xlsx header cell holding the number 1
    # arrives from pandas as the int 1, which would never equal a requested
    # `target="1"` and would mix types in every name list downstream.
    parsed.df.columns = parsed.df.columns.map(str)
    return parsed


def _parse_upload(raw: bytes) -> pd.DataFrame:
    """Backward-compatible wrapper: `parse_upload(raw).df` (first usable sheet
    of a workbook; no filename-based checks)."""
    return parse_upload(raw).df
